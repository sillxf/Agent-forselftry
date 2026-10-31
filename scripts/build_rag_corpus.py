"""Build the small v1.1 corpus. Install requirements-corpus.txt; optionally use --download."""
import argparse
import hashlib
import io
import json
import tempfile
import urllib.request
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import Table, TableStyle
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1] / "tests/fixtures/rag_documents"
COMMIT = "764b5d4dc22b03ea87afd4c2207f510300b27a5b"
NIST_URL = "https://nvlpubs.nist.gov/nistpubs/SpecialPublications/NIST.SP.800-61r3.pdf"
DOCLING_FILES = {
    "xlsx/xlsx_01.xlsx": "xlsx/sources/xlsx_01.xlsx",
    "xlsx/xlsx_05_table_with_title.xlsx": "xlsx/sources/xlsx_05_table_with_title.xlsx",
    "pptx/powerpoint_sample.pptx": "pptx/sources/powerpoint_sample.pptx",
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def download(url, target, proxy):
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({"https": proxy})) if proxy else urllib.request.build_opener()
    with opener.open(url, timeout=90) as response:
        target.write_bytes(response.read())


def stable_docx(document, path):
    document.save(path)
    stream = io.BytesIO()
    with ZipFile(path) as original, ZipFile(stream, "w", ZIP_DEFLATED) as result:
        for info in original.infolist():
            info.date_time = (2026, 10, 31, 0, 0, 0)
            result.writestr(info, original.read(info.filename))
    path.write_bytes(stream.getvalue())


def paragraph(document, text, size=11, bold=False):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.font.name, run.font.size, run.bold = "Arial", Pt(size), bold
    return p


def pdf_page(c, page):
    c.setFont("Helvetica", 9)
    c.drawString(54, 760, "SOC validation manual")
    c.drawRightString(558, 32, str(page))
    c.setFont("Helvetica", 12)


def main():
    arguments = argparse.ArgumentParser(description=__doc__)
    arguments.add_argument("--download", action="store_true", help="Fetch pinned Docling files and the NIST source")
    arguments.add_argument("--proxy")
    args = arguments.parse_args()
    for folder in ("pdf/text_based", "pdf/scanned", "pdf/tables", "pdf/cross_page", "docx", "xlsx", "pptx", "images"):
        (ROOT / folder).mkdir(parents=True, exist_ok=True)
    nist_source = Path(tempfile.gettempdir()) / "secops-nist-sp80061r3.pdf"
    if args.download:
        download(NIST_URL, nist_source, args.proxy)
        for target, source in DOCLING_FILES.items():
            download(f"https://raw.githubusercontent.com/docling-project/docling/{COMMIT}/tests/data/{source}", ROOT / target, args.proxy)
        download(f"https://raw.githubusercontent.com/docling-project/docling/{COMMIT}/LICENSE", ROOT / "DOCLING_LICENSE.txt", args.proxy)
    entries = []

    def record(file, purpose, expected, origin=None):
        entries.append({"file": file, "format": Path(file).suffix[1:], "sha256": digest(ROOT / file),
                        "origin": origin or {"kind": "generated", "generator": "scripts/build_rag_corpus.py"},
                        "requires_models": Path(file).suffix in {".pdf", ".png"},
                        "purpose": purpose, "expected": expected})

    # One real security document, a clearly identified excerpt rather than a fabricated guide.
    nist = PdfReader(nist_source)
    writer = PdfWriter()
    for page in nist.pages[9:12]:
        writer.add_page(page)
    file = "pdf/text_based/nist_sp80061r3_pp10-12.pdf"
    writer.write(ROOT / file)
    record(file, ["real_security_document", "headings", "page_provenance"],
           {"contains": ["Introduction", "Cybersecurity Risk Management"], "pages": [1, 2, 3]},
           {"kind": "official_excerpt", "url": NIST_URL, "version": "SP 800-61 Rev. 3, April 2025",
            "source_sha256": digest(nist_source), "physical_source_pages": [10, 11, 12],
            "license": "US NIST government publication; original attribution retained"})

    # Real file boundaries, including a sentence deliberately broken across two pages.
    file = "pdf/cross_page/sentence_continuation.pdf"
    c = canvas.Canvas(str(ROOT / file), pagesize=letter, invariant=1)
    pdf_page(c, 1); c.setFont("Helvetica-Bold", 17); c.drawString(54, 710, "Authentication investigation")
    c.setFont("Helvetica", 12); c.drawString(54, 100, "The authentication rule checks events,"); c.showPage()
    pdf_page(c, 2); c.drawString(54, 680, "then correlates failed login attempts across accounts."); c.save()
    record(file, ["cross_page_sentence", "page_furniture"],
           {"merged_text": "The authentication rule checks events, then correlates failed login attempts across accounts.",
            "merged_pages": [1, 2], "excludes": ["SOC validation manual"]})

    file = "pdf/tables/continued_rules.pdf"
    c = canvas.Canvas(str(ROOT / file), pagesize=letter, invariant=1)
    for page, rows in [(1, [["R001", "High", "Investigate"], ["R002", "Critical", "Isolate"]]),
                       (2, [["R003", "Medium", "Review"], ["R004", "High", "Block"]])]:
        pdf_page(c, page); c.setFont("Helvetica-Bold", 17)
        c.drawString(54, 710, "Detection Rules" + (" (continued)" if page == 2 else ""))
        table = Table([["Rule ID", "Severity", "Response"], *rows], colWidths=[110, 130, 190], rowHeights=32)
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .7, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                                   ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                                   ("FONTSIZE", (0, 0), (-1, -1), 12)]))
        table.wrapOn(c, 450, 200); table.drawOn(c, 54, 540); c.showPage()
    c.save()
    record(file, ["repeated_table_header", "cross_page_table"],
           {"table_values": ["R001", "R002", "R003", "R004"], "table_count": 1,
            "merged_pages": [1, 2], "merged_type": "table"})

    font = ImageFont.load_default(size=42)
    for name, lines in [("no_ocr", []), ("ocr_note", ["IR RESPONSE NOTE", "COLLECT LOGS"]),
                        ("multi_region_ocr", ["COLLECT LOGS", "ISOLATE HOST", "BLOCK SOURCE IP"])]:
        image = Image.new("RGB", (1200, 1000), "white")
        draw = ImageDraw.Draw(image)
        if not lines:
            draw.rectangle((150, 250, 500, 600), outline="navy", width=10)
            draw.ellipse((700, 300, 1000, 600), outline="green", width=10)
        for index, line in enumerate(lines):
            draw.text((90, 150 + index * 250), line, font=font, fill="black")
        file = f"images/{name}.png"; image.save(ROOT / file)
        expected = {"source_image": True}
        if lines:
            expected["contains"] = lines
        record(file, ["source_image_metadata", "no_ocr" if not lines else "ocr_regions"], expected)

    file = "pdf/scanned/incident_note.pdf"
    c = canvas.Canvas(str(ROOT / file), pagesize=letter, invariant=1)
    c.drawImage(str(ROOT / "images/ocr_note.png"), 0, 140, width=612, height=510)
    c.save()
    record(file, ["image_only_pdf", "ocr"], {"contains": ["IR RESPONSE NOTE", "COLLECT LOGS"], "pages": [1]})

    file = "pdf/cross_page/figure_interrupts_text.pdf"
    c = canvas.Canvas(str(ROOT / file), pagesize=letter, invariant=1)
    pdf_page(c, 1); c.setFont("Helvetica-Bold", 17); c.drawString(54, 710, "Network investigation")
    c.setFont("Helvetica", 12); c.drawString(54, 100, "The network investigation continues,"); c.showPage()
    pdf_page(c, 2); c.drawImage(str(ROOT / "images/no_ocr.png"), 70, 450, width=240, height=200)
    c.drawString(54, 350, "then follows the illustrated connection path."); c.save()
    record(file, ["picture_between_page_continuations", "known_limitation"],
           {"contains": ["network investigation continues", "illustrated connection path"], "picture_min_count": 1})
    entries[-1]["known_limitation"] = "Page-level candidate matching is deferred: a picture can interrupt paragraph merging."

    file = "docx/duplicate_paragraphs.docx"
    doc = Document(); paragraph(doc, "Detection", 18, True)
    paragraph(doc, "The first occurrence introduces the incident detection topic.")
    paragraph(doc, "Detection", 10, False); paragraph(doc, "The second occurrence is ordinary source text.")
    stable_docx(doc, ROOT / file)
    record(file, ["duplicate_text_layout"], {"layout_occurrences": [{"text": "Detection", "font_sizes": [18, 10], "bold": [True, False]}],
                                             "contains": ["first occurrence", "second occurrence"]})

    file = "docx/numbered_steps.docx"
    doc = Document(); doc.add_heading("Response checklist", 0)
    steps = ["1. Check firewall configuration", "2. Collect authentication logs", "3. Isolate affected hosts",
             "1. 检查防火墙配置", "2. 收集认证日志", "3. 隔离受影响主机"]
    for text in steps:
        paragraph(doc, text)
    stable_docx(doc, ROOT / file)
    record(file, ["heading_false_positive", "numbered_steps"], {"not_headings": steps, "contains": ["firewall configuration", "收集认证日志"]})

    file = "docx/no_heading.docx"
    doc = Document(); paragraph(doc, "Detection", 16, True)
    paragraph(doc, "Repeated failed login attempts require investigation of affected accounts and authentication logs.")
    paragraph(doc, "Response", 16, True)
    paragraph(doc, "Isolate affected hosts and collect forensic evidence before restoring the service.")
    stable_docx(doc, ROOT / file)
    record(file, ["layout_headings", "no_native_headings"], {"inferred_headings": ["Detection", "Response"], "contains": ["failed login attempts", "forensic evidence"]})

    file = "docx/cross_page_lists.docx"
    doc = Document(); doc.add_heading("Independent checklist A", 1)
    doc.add_paragraph("Inspect the authentication logs", "List Bullet")
    doc.add_page_break(); doc.add_heading("Independent checklist B", 1)
    doc.add_paragraph("Review the network events", "List Bullet")
    stable_docx(doc, ROOT / file)
    record(file, ["independent_lists", "explicit_page_break"], {"contains": ["authentication logs", "network events"], "merge_count": 0})

    for file, source in DOCLING_FILES.items():
        expected = {"table_min_count": 6, "contains": ["first", "second", "col-1"]} if file.endswith("xlsx_01.xlsx") else (
            {"table_min_count": 2, "contains": ["Freshwater Ducks", "2024"]} if file.endswith("xlsx_05_table_with_title.xlsx") else
            {"table_min_count": 1, "pages": [1, 2, 3]})
        record(file, ["upstream_real_file", "tables_and_merged_cells" if file.startswith("xlsx") else "slide_layout"], expected,
               {"kind": "upstream_fixture", "url": f"https://raw.githubusercontent.com/docling-project/docling/{COMMIT}/tests/data/{source}",
                "commit": COMMIT, "license": "MIT; see DOCLING_LICENSE.txt"})
    manifest = {"version": 1, "documents": entries}
    (ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Built {len(entries)} files with {sum(len(e['expected']) for e in entries)} expected-behaviour groups")


if __name__ == "__main__":
    main()
