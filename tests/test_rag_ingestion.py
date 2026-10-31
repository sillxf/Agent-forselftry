from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pandas as pd
import pytest
from PIL import Image

from app.rag.ingestion.docling_parser import DoclingParser
from app.rag.loader import load_documents


def test_markdown_headings_tables_and_archives(tmp_path):
    path = tmp_path / "manual.md"
    path.write_text("# Security\n\n## Detection\n\nText evidence\n\n### Rules\n\n"
                    "| Rule ID | Severity |\n| --- | --- |\n| R001 | High |\n", encoding="utf-8")
    parser = DoclingParser(archive_dir=tmp_path / "parsed", image_dir=tmp_path / "images")
    parsed = parser.parse(path)
    table = next(b for b in parsed.blocks if b.block_type == "table")
    assert table.table_data == [{"Rule ID": "R001", "Severity": "High"}]
    assert table.section_path == ["Security", "Detection", "Rules"]
    assert table.page_number is None
    assert "Severity: High" in table.text
    assert (tmp_path / "parsed" / f"{parsed.document.checksum}.json").exists()
    assert (tmp_path / "parsed" / f"{parsed.document.checksum}.md").exists()
    assert parser.parse(path).blocks[0].block_id == parsed.blocks[0].block_id
    assert load_documents(tmp_path)[0]["filename"] == "manual.md"


@pytest.mark.parametrize("extension", ["docx", "xlsx", "pptx", "html"])
def test_actual_docling_office_and_html_conversion(tmp_path, extension):
    path = tmp_path / f"rules.{extension}"
    if extension == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading("Detection", level=1)
        doc.add_paragraph("Rule R001 evidence")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Rule ID", "Severity"
        table.cell(1, 0).text, table.cell(1, 1).text = "R001", "High"
        doc.save(path)
    elif extension == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "Detection"
        sheet.append(["Rule ID", "Severity"])
        sheet.append(["R001", "High"])
        book.save(path)
    elif extension == "pptx":
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Detection"
        slide.placeholders[1].text = "Rule R001 evidence"
        deck.save(path)
    else:
        path.write_text("<html><body><h1>Detection</h1><p>Rule R001 evidence</p>"
                        "<table><tr><th>Rule ID</th><th>Severity</th></tr>"
                        "<tr><td>R001</td><td>High</td></tr></table></body></html>", encoding="utf-8")
    parsed = DoclingParser().parse(path)
    assert parsed.document.file_type == extension
    assert "R001" in "\n".join(b.text for b in parsed.blocks)
    if extension != "pptx":
        assert any(b.block_type == "table" and b.table_data for b in parsed.blocks)


def test_pdf_provenance_tables_and_picture_links_via_converter_contract(tmp_path):
    path = tmp_path / "manual.pdf"
    path.write_bytes(b"test PDF converter contract")

    def item(label, ref, pages, **kwargs):
        return SimpleNamespace(label=label, self_ref=ref, prov=[SimpleNamespace(page_no=p) for p in pages], **kwargs)

    items = [item("section_header", "h", [4], text="Detection", level=1),
             item("text", "p1", [4], text="Investigation starts,"),
             item("text", "p2", [5], text="continues on next page."),
             item("table", "t", [5, 6], export_to_dataframe=lambda doc: pd.DataFrame([["R001", "High"]], columns=["Rule ID", "Severity"]),
                  caption_text=lambda doc: "Rules"),
             item("picture", "i", [6], children=[SimpleNamespace(resolve=lambda doc: SimpleNamespace(text="Firewall Web Server"))],
                  caption_text=lambda doc: "Architecture",
                  get_image=lambda doc: Image.new("RGB", (4, 4)))]
    doc = SimpleNamespace(iterate_items=lambda **kwargs: [(i, 1) for i in items],
                          export_to_markdown=lambda: "Detection", export_to_dict=lambda: {"name": "manual"})
    converter = Mock()
    converter.convert.return_value = SimpleNamespace(status="success", document=doc)
    parsed = DoclingParser(converter=converter, image_dir=tmp_path / "images").parse(path)
    table, image = parsed.blocks[-2:]
    assert table.page_number == 5 and table.page_end == 6
    assert table.table_data[0]["Severity"] == "High"
    assert image.page_number == 6
    assert "Firewall Web Server" in image.text
    assert image.image_path and image.related_block_ids == [table.block_id]
    assert Image.open(image.image_path).size == (4, 4)
    assert parsed.blocks[2].section_path == ["Detection"]

    from app.rag.chunking.child_chunker import ChildChunker
    from app.rag.chunking.parent_chunker import ParentChunker
    from app.rag.processing.cross_page_merger import CrossPageMerger
    from app.rag.processing.structure_recovery import StructureRecovery
    from conftest import CharacterTokenizer

    blocks = StructureRecovery().recover(CrossPageMerger().merge(parsed.blocks))
    parents = ParentChunker(CharacterTokenizer()).split(blocks, parsed.document)
    children = [child for parent in parents for child in ChildChunker(CharacterTokenizer(), 240, 20).split(parent, path.name)]
    evidence = next(child for child in children if "Investigation starts," in child.content)
    assert "continues on next page." in evidence.content
    assert set([parsed.blocks[1].block_id, parsed.blocks[2].block_id]) <= set(evidence.block_ids)
    assert evidence.page_start == 4 and evidence.page_end == 5


@pytest.mark.parametrize("texts", [[], ["COLLECT LOGS"], ["COLLECT LOGS", "ISOLATE HOST", "BLOCK SOURCE IP"]])
def test_image_source_metadata_is_independent_of_ocr_items(tmp_path, texts):
    path = Path(__file__).parent / "fixtures/rag_documents/images/no_ocr.png"
    items = [SimpleNamespace(label="section_header", self_ref=f"ocr-{index}", text=text, level=3,
                              prov=[SimpleNamespace(page_no=1)], formatting=SimpleNamespace(bold=True))
             for index, text in enumerate(texts)]
    converter = Mock()
    converter.convert.return_value = SimpleNamespace(status="success", document=SimpleNamespace(
        iterate_items=lambda **kwargs: [(item, 1) for item in items],
        export_to_markdown=lambda: "\n".join(texts), export_to_dict=lambda: {}))
    parsed = DoclingParser(converter=converter, image_dir=tmp_path).parse(path)
    source = parsed.blocks[0]
    assert source.block_type == "image" and source.image_path
    assert source.caption is None and source.heading_level is None and source.list_marker is None
    assert not source.layout.bold and source.layout.font_size is None
    assert source.structure_source == "unknown"
    assert source.related_block_ids == [block.block_id for block in parsed.blocks[1:]]
    assert all(text in source.text for text in texts)


def test_failed_image_conversion_does_not_create_a_source_record(tmp_path):
    path = Path(__file__).parent / "fixtures/rag_documents/images/no_ocr.png"
    converter = Mock()
    converter.convert.return_value = SimpleNamespace(status="failure")
    with pytest.raises(ValueError, match="conversion failed"):
        DoclingParser(converter=converter, image_dir=tmp_path).parse(path)
    assert not list(tmp_path.iterdir())
