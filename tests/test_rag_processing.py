"""Representative continuation/structure cases, plus the real ingestion-to-search path."""
from app.rag.models import BlockLayout, ContentBlock
from app.rag.processing.cross_page_merger import CrossPageMerger
from app.rag.processing.structure_recovery import StructureRecovery


def block(text, page=1, kind="paragraph", section=None, **metadata):
    return ContentBlock(document_id="manual", text=text, block_type=kind,
                        page_number=page, section_path=section or [], **metadata)


def test_paragraph_continuation_removes_page_noise_and_keeps_sources():
    first = block("The detector checks authenti-", 1)
    second = block("cation failures.", 2)
    blocks = [block("SOC manual", 1), first, block("1", 1),
              block("SOC manual", 2), second, block("2", 2)]
    result = CrossPageMerger().merge(blocks)
    assert len(result) == 1
    assert result[0].text == "The detector checks authentication failures."
    assert result[0].page_number == 1 and result[0].page_end == 2
    assert result[0].source_block_ids == [first.block_id, second.block_id]
    assert first.text.endswith("authenti-")  # Original blocks remain unchanged.


def test_table_headers_rows_and_section_boundaries():
    first = block("Rules", 4, "table", ["Detection"], caption="Rules",
                  table_columns=["Rule ID", "Description"], table_data=[{"Rule ID": "R001", "Description": "Failed login,"}])
    second = block("Rules continued", 5, "table", ["Detection"], caption="Rules (continued)",
                   table_columns=["Column 1", "Column 2"],
                   table_data=[{"Column 1": "Rule ID", "Column 2": "Description"},
                               {"Column 1": "", "Column 2": "over five minutes"},
                               {"Column 1": "R002", "Column 2": "SQL injection"}])
    third = block("Other rules", 6, "table", ["Response"], caption="Response rules",
                  table_columns=first.table_columns, table_data=[{"Rule ID": "R003", "Description": "Isolate asset"}])
    result = CrossPageMerger().merge([first, second, third])
    assert len(result) == 2
    assert result[0].table_data == [{"Rule ID": "R001", "Description": "Failed login, over five minutes"},
                                    {"Rule ID": "R002", "Description": "SQL injection"}]
    assert result[0].page_end == 5
    assert result[1].section_path == ["Response"]


def test_lists_continue_but_restarts_and_new_topics_do_not_merge():
    items = [block("Inspect the logs", 1, "list", list_marker="1."),
             block("Isolate affected hosts", 2, "list", list_marker="2."),
             block("Start a different checklist", 3, "list", list_marker="1.")]
    result = CrossPageMerger().merge(items)
    assert len(result) == 2
    assert result[0].text == "1. Inspect the logs\n2. Isolate affected hosts"
    assert result[0].page_end == 2
    # A heading-like paragraph is a boundary even before structure recovery.
    paragraphs = [block("Investigation continues,", 1), block("2. Response", 2), block("isolate affected hosts", 2)]
    assert len(CrossPageMerger().merge(paragraphs)) == 3


def test_numbered_and_layout_headings_are_inferred_without_replacing_originals():
    blocks = [block("SOC manual", kind="heading", section=["SOC manual"], heading_level=0),
              block("1. Detection"), block("Inspect authentication events."),
              block("1.1 Evidence"), block("Collect source IPs and account names."),
              block("2. Response"), block("Isolate affected assets.")]
    recovered = StructureRecovery().recover(blocks)
    assert recovered[0].structure_source == "original"
    assert recovered[4].section_path == ["SOC manual", "1. Detection", "1.1 Evidence"]
    assert recovered[-1].section_path == ["SOC manual", "2. Response"]
    assert recovered[1].structure_source == "inferred" and recovered[1].structure_reason == "numbering"
    layout_blocks = [block("Incident handling", layout=BlockLayout(font_size=20, bold=True)),
                     block("A long body paragraph. " * 5, layout=BlockLayout(font_size=10)),
                     block("Evidence sources", layout=BlockLayout(font_size=14, bold=True)),
                     block("Another long body paragraph. " * 5, layout=BlockLayout(font_size=10))]
    layout_result = StructureRecovery().recover(layout_blocks)
    assert layout_result[-1].section_path == ["Incident handling", "Evidence sources"]
    assert layout_result[2].structure_reason == "font_size"
    native = StructureRecovery().recover([
        block("Native chapter", kind="heading", section=["Native chapter"], heading_level=1),
        block("1. Detection"), block("1.1 Evidence"), block("Record the relevant events."),
    ])
    assert native[-1].section_path == ["Native chapter", "1. Detection", "1.1 Evidence"]


def test_unheaded_topic_groups_require_context_and_ambiguous_text_stays_unknown():
    blocks = [block("检测告警并收集认证日志，调查来源 IP。"), block("排查日志中的异常，确认告警涉及的账户。"),
              block("处置时隔离受影响的主机并封禁来源。"), block("修复服务后恢复运行，完成处置记录。")]
    recovered = StructureRecovery().recover(blocks)
    assert [b.section_path for b in recovered] == [["检测与调查"], ["检测与调查"], ["处置与修复"], ["处置与修复"]]
    assert all(b.structure_source == "inferred" for b in recovered)
    ambiguous = StructureRecovery().recover([block("Some standalone information.")])[0]
    assert ambiguous.section_path == [] and ambiguous.structure_source == "unknown"
    entities = StructureRecovery().recover([
        block("CVE-2025-12345: " + "The affected service exposes an exploitable endpoint. " * 3),
        block("CVE-2025-67890: " + "The affected asset requires an independent assessment. " * 3),
    ])
    assert [b.section_path for b in entities] == [["CVE-2025-12345"], ["CVE-2025-67890"]]
    mixed = StructureRecovery().recover([block("CVE-2025-12345: Summary"), entities[1]])
    assert mixed[1].section_path == ["CVE-2025-67890"]


def test_docx_layout_recovery_is_used_by_retrieval(rag_factory):
    from docx import Document
    from docx.shared import Pt

    path = rag_factory.knowledge / "unheaded.docx"
    document = Document()
    title = document.add_paragraph()
    run = title.add_run("Authentication investigation")
    run.bold, run.font.size = True, Pt(18)
    body = document.add_paragraph("Repeated failed login attempts require investigation of affected accounts.")
    body.runs[0].font.size = Pt(11)
    document.save(path)
    retriever = rag_factory()
    results = retriever.search("failed login attempts", 1)
    assert "failed login attempts" in results[0]["content"]
    assert results[0]["section_path"] == ["Authentication investigation"]
    assert results[0]["structure_source"] == "inferred"
    assert results[0]["source"] == "unheaded.docx"
